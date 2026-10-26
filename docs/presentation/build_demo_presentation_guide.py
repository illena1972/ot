from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = OUTPUT_DIR / "Сценарий презентации BioClean Workwear.docx"

BLACK = "000000"
DARK_BLUE = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF4FA"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "666666"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char1, instr_text, fld_char2))


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, before, after in (
        ("Title", 25, 0, 10),
        ("Heading 1", 17, 12, 7),
        ("Heading 2", 13, 9, 5),
        ("Heading 3", 11, 7, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "BioClean Workwear | Сценарий презентации"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    add_page_number(section.footer.paragraphs[0])


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, title in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[index])

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(str(value))
            run.font.name = "Arial"
            run.font.size = Pt(9)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_script(doc, lead, text):
    paragraph = doc.add_paragraph()
    lead_run = paragraph.add_run(f"{lead} ")
    lead_run.bold = True
    paragraph.add_run(text)
    return paragraph


def page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Сценарий презентации BioClean Workwear")
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Как провести встречу на демонстрационной базе и выявить потребности клиента")
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    doc.add_paragraph()
    meta = add_table(
        doc,
        ("Параметр", "Значение"),
        (
            ("Рекомендуемая длительность", "25–35 минут"),
            ("Демонстрационный адрес", "http://demo.localhost:5173"),
            ("Логин", "demo_admin"),
            ("Версия сценария", "26 октября 2026 г."),
        ),
        widths=(2.1, 4.5),
    )
    meta.rows[0].cells[0].width = Inches(2.1)

    doc.add_heading("Главный принцип", level=1)
    doc.add_paragraph(
        "Презентация должна начинаться не с перечня функций, а с рабочего процесса клиента. "
        "Сначала выясните, как сейчас контролируются нормы, выдачи и сроки, затем показывайте только те возможности, "
        "которые отвечают на названные проблемы. Цель встречи — не показать каждый экран, а подтвердить, что система "
        "поможет ответственному сотруднику вовремя увидеть отклонение и выполнить следующую операцию."
    )

    doc.add_heading("Результат встречи", level=2)
    add_bullet(doc, "Клиент понимает путь: должность → норма → выдача → контроль срока → заказ.")
    add_bullet(doc, "Вы фиксируете масштаб, обязательные процессы и ограничения клиента.")
    add_bullet(doc, "Обе стороны договариваются о следующем шаге: тестовый доступ, уточнение данных или коммерческое предложение.")

    page_break(doc)
    doc.add_heading("1. Подготовка перед встречей", level=1)
    doc.add_paragraph("Проверьте систему за 15–20 минут до подключения клиента.")
    add_table(
        doc,
        ("Проверка", "Что должно быть готово"),
        (
            ("Доступ", "Открывается demo.localhost:5173, вход под demo_admin выполняется без ошибок."),
            ("Рабочий экран", "Видны три группы: просрочено, срок заканчивается, не выдано."),
            ("Нормы", "Отображаются три нормы, должности и развернутый состав СИЗ."),
            ("Сотрудники", "Есть Иванов, Петрова, Сидоров, Орлова и Кузнецов."),
            ("Склад", "Есть остатки по верхней одежде, обуви и безразмерным СИЗ."),
            ("Выдача", "Форма открывается, позиции по норме можно добавить в документ."),
            ("Связь", "Закрыты лишние окна и уведомления, демонстрация экрана работает."),
        ),
        widths=(1.45, 5.15),
    )

    doc.add_heading("Обновление сценарных дат", level=2)
    doc.add_paragraph(
        "Если демонстрационная база использовалась спустя длительное время после создания, пересоздайте только её данные. "
        "Логин администратора сохранится, а сроки снова будут рассчитаны относительно текущей даты."
    )
    command = doc.add_paragraph()
    command.paragraph_format.left_indent = Inches(0.25)
    run = command.add_run(
        ".\\venv\\Scripts\\python.exe manage.py seed_demo_organization "
        "--organization demo --reset --confirm-reset-demo"
    )
    run.font.name = "Consolas"
    run.font.size = Pt(9)

    doc.add_heading("Что не показывать", level=2)
    add_bullet(doc, "Рабочую базу First и реальные ФИО сотрудников.")
    add_bullet(doc, "Файл .env, пароли, настройки базы данных и служебные журналы.")
    add_bullet(doc, "Незавершенные функции или обещания, которых нет в согласованной версии продукта.")

    page_break(doc)
    doc.add_heading("2. Рекомендуемая структура встречи", level=1)
    add_table(
        doc,
        ("Время", "Этап", "Цель"),
        (
            ("0–3 мин", "Контекст", "Представиться, согласовать цель и длительность встречи."),
            ("3–8 мин", "Вопросы", "Понять текущий процесс, масштаб и основные затруднения."),
            ("8–12 мин", "Контроль СИЗ", "Показать единый список задач и предупреждение по срокам."),
            ("12–17 мин", "Нормы и обеспеченность", "Показать связь должности, нормы и фактической выдачи."),
            ("17–22 мин", "Выдача и склад", "Показать рабочую операцию и автоматическое изменение остатка."),
            ("22–26 мин", "Отчёты", "Показать подготовку данных для контроля и заказа."),
            ("26–35 мин", "Обсуждение", "Уточнить соответствие процессу клиента и согласовать следующий шаг."),
        ),
        widths=(0.8, 1.65, 4.15),
    )

    doc.add_heading("Начало разговора", level=2)
    add_script(
        doc,
        "Скажите:",
        "«Предлагаю сначала коротко понять, как у вас сейчас организован учёт СИЗ, а затем показать систему на похожих задачах. "
        "Так мы не будем тратить время на функции, которые для вас несущественны»."
    )
    add_script(
        doc,
        "Уточните:",
        "«Правильно ли я понимаю, что основная цель — заранее видеть истекающие сроки и понимать, кому и что необходимо выдать?»"
    )

    page_break(doc)
    doc.add_heading("3. Вопросы до демонстрации", level=1)
    doc.add_paragraph(
        "Не задавайте весь список подряд. Выберите 6–8 вопросов, а остальные используйте для уточнения по ходу разговора."
    )

    question_groups = (
        ("Текущий процесс", (
            "Как сейчас ведётся учёт: Excel, бумажные карточки, 1С или другая система?",
            "Кто отвечает за нормы, выдачу, склад и контроль сроков?",
            "Какая ошибка или задержка сейчас создаёт наибольший риск?",
            "Как часто руководителю требуется сводная информация?",
        )),
        ("Масштаб", (
            "Сколько сотрудников, подразделений и должностей участвует в учёте?",
            "Один ли склад и одна организация должны работать в системе?",
            "Сколько пользователей будут оформлять операции одновременно?",
        )),
        ("Нормы и сроки", (
            "Утверждены ли нормы по должностям и в каком виде они хранятся?",
            "Есть ли индивидуальные исключения из норм?",
            "За сколько дней до окончания срока нужно предупреждать: 30, 60, 90, 180 или 360?",
            "Что важнее видеть первым: просрочку, ближайшую замену или невыданные позиции?",
        )),
        ("Выдача и склад", (
            "Кто оформляет выдачу и где происходит фактическая передача СИЗ?",
            "Нужно ли учитывать размеры и рост, а также безразмерные позиции?",
            "Как оформляется списание и требуется ли возврат на склад?",
            "Есть ли несколько складов или материально ответственных лиц?",
        )),
        ("Внедрение", (
            "Предпочтительна работа в облаке или внутри локальной сети?",
            "Какие данные нужно подготовить к началу работы?",
            "Нужны ли разные права доступа для пользователей?",
            "Какие интеграции обязательны для запуска?",
        )),
    )
    for heading, questions in question_groups:
        doc.add_heading(heading, level=2)
        for question in questions:
            add_bullet(doc, question)

    page_break(doc)
    doc.add_heading("4. Маршрут показа по демонстрационной базе", level=1)

    doc.add_heading("Шаг 1. Контроль СИЗ", level=2)
    add_script(
        doc,
        "Скажите:",
        "«Это постоянный рабочий экран специалиста. Здесь не нужно просматривать карточку каждого сотрудника: система сама собирает задачи по срокам и нормам»."
    )
    add_number(doc, "Покажите переключатель периода предупреждения и значение 60 дней.")
    add_number(doc, "Последовательно выберите фильтры «Просрочено», «Срок заканчивается» и «Не выдано».")
    add_number(doc, "Обратите внимание на сортировку: сначала статус, затем ФИО сотрудника.")

    add_table(
        doc,
        ("Сотрудник", "Что демонстрирует"),
        (
            ("Иванов Алексей Сергеевич", "Просроченная каска и ботинки с приближающимся сроком замены."),
            ("Петрова Марина Викторовна", "Скорую замену халата и позиции, которые ещё не выданы по норме."),
            ("Сидоров Николай Андреевич", "Полностью не обеспеченную норму нового сотрудника."),
            ("Орлова Елена Игоревна", "Сотрудника, обеспеченного действующими СИЗ."),
            ("Кузнецов Павел Олегович", "Контроль фактической выдачи даже при отсутствии нормы у должности."),
        ),
        widths=(2.35, 4.25),
    )

    doc.add_heading("Шаг 2. Нормы выдачи", level=2)
    add_script(
        doc,
        "Скажите:",
        "«Норма назначается должности и содержит перечень СИЗ, количество и срок эксплуатации. Одну норму можно связать с несколькими должностями»."
    )
    add_number(doc, "Откройте список норм и разверните состав одной нормы.")
    add_number(doc, "Покажите поиск по названию нормы или должности.")
    add_number(doc, "Откройте редактирование, но не меняйте данные во время основной презентации.")

    page_break(doc)
    doc.add_heading("5. Продолжение маршрута", level=1)

    doc.add_heading("Шаг 3. Обеспеченность СИЗ", level=2)
    add_script(
        doc,
        "Скажите:",
        "«В карточке обеспеченности видно, что положено сотруднику, что уже выдано, чего не хватает и когда наступит следующая выдача»."
    )
    add_number(doc, "Выберите Петрову и покажите одновременно действующие, заканчивающиеся и невыданные позиции.")
    add_number(doc, "Выберите Орлову как пример полностью обеспеченного сотрудника.")
    add_number(doc, "Подчеркните, что срок берётся из нормы при оформлении новой выдачи.")

    doc.add_heading("Шаг 4. Выдача одежды", level=2)
    add_script(
        doc,
        "Скажите:",
        "«Специалист выбирает сотрудника, видит положенные позиции, добавляет нужные строки и оформляет один документ выдачи»."
    )
    add_number(doc, "Выберите Сидорова, чтобы показать подсказки по невыданной норме.")
    add_number(doc, "Добавьте одну позицию и покажите автоматически подставленный срок.")
    add_number(doc, "Не завершайте выдачу без необходимости; для живой демонстрации заранее выберите безопасную позицию с достаточным остатком.")

    doc.add_heading("Шаг 5. Остатки и отчёты", level=2)
    add_script(
        doc,
        "Скажите:",
        "«Склад ведётся по наименованию, размеру и росту. После выдачи остаток уменьшается. Отчёт для заказа объединяет потребность и текущий запас»."
    )
    add_number(doc, "Покажите фильтры склада по типу одежды.")
    add_number(doc, "Откройте отчёт по выдаче и отчёт для заказа.")
    add_number(doc, "Покажите экспорт, если он важен для процесса клиента.")

    page_break(doc)
    doc.add_heading("6. Вопросы во время показа", level=1)
    add_table(
        doc,
        ("После экрана", "Вопрос клиенту"),
        (
            ("Контроль СИЗ", "Такой список задач заменил бы ваш текущий контроль сроков? Чего в нём не хватает?"),
            ("Нормы", "Насколько ваши нормы похожи на эту структуру? Есть ли сезонность или индивидуальные исключения?"),
            ("Обеспеченность", "Кто должен видеть эту информацию и как часто?"),
            ("Выдача", "Совпадает ли эта последовательность с вашей фактической процедурой?"),
            ("Склад", "Достаточно одного общего остатка или нужно разделение по складам?"),
            ("Отчёт для заказа", "По каким правилам вы сейчас формируете заявку поставщику?"),
        ),
        widths=(1.55, 5.05),
    )

    doc.add_heading("Как фиксировать ответы", level=2)
    doc.add_paragraph("После каждого существенного ответа помечайте требование одной из трёх категорий:")
    add_bullet(doc, "Обязательно для запуска — без этого клиент не сможет начать работу.")
    add_bullet(doc, "Желательно — повышает удобство, но не блокирует запуск.")
    add_bullet(doc, "Перспектива — имеет смысл после проверки базового процесса.")

    doc.add_heading("Признаки реальной заинтересованности", level=2)
    add_bullet(doc, "Клиент уточняет сроки внедрения, состав исходных данных и ответственных лиц.")
    add_bullet(doc, "Просит показать конкретную операцию на своём примере.")
    add_bullet(doc, "Обсуждает число пользователей, подразделений или порядок согласования покупки.")
    add_bullet(doc, "Готов предоставить обезличенный пример нормы или отчёта для проверки соответствия.")

    page_break(doc)
    doc.add_heading("7. Честные ответы на ожидаемые вопросы", level=1)
    add_table(
        doc,
        ("Вопрос", "Рекомендуемый ответ"),
        (
            ("Есть ли уведомления по электронной почте?", "Сейчас предупреждения работают внутри программы на экране контроля СИЗ. Внешние уведомления можно обсуждать как отдельное развитие."),
            ("Можно ли выбрать срок предупреждения?", "Да. Для организации выбирается один из периодов: 30, 60, 90, 180 или 360 дней."),
            ("Что будет с сотрудником без нормы?", "Система контролирует сроки уже выданных ему СИЗ. Статус «Не выдано» рассчитывается только при наличии нормы."),
            ("Есть ли разные роли и ограничения доступа?", "Разделение прав по ролям пока не входит в текущую пользовательскую версию. Требования к ролям нужно зафиксировать отдельно."),
            ("Можно ли работать локально?", "Да. Поддерживаются локальная установка в сети организации и SaaS-версия на хостинге."),
            ("Как защищены данные разных клиентов?", "В SaaS каждая организация работает со своей отдельной базой данных и собственным адресом."),
            ("Как выполняется резервное копирование?", "Для серверной и локальной установки предусмотрены регулярные резервные копии баз. Конкретный режим подтверждается при внедрении."),
            ("Переносятся ли данные автоматически?", "Состав и способ начальной загрузки согласуются после анализа исходных файлов. Не обещайте автоматический перенос до проверки данных."),
            ("Сколько стоит система?", "Используйте актуальное коммерческое предложение или сайт; не называйте стоимость по памяти."),
        ),
        widths=(2.2, 4.4),
    )

    doc.add_heading("Чего не обещать", level=2)
    add_bullet(doc, "Функции, которых ещё нет в текущей версии.")
    add_bullet(doc, "Бесплатный или полностью автоматический перенос любых таблиц Excel.")
    add_bullet(doc, "Срок разработки индивидуальной функции до оценки задачи.")
    add_bullet(doc, "Интеграцию с внешней системой без изучения её API и форматов.")

    page_break(doc)
    doc.add_heading("8. Завершение встречи", level=1)
    add_script(
        doc,
        "Подведите итог:",
        "«Сегодня мы увидели, как система связывает должность, норму, фактическую выдачу и контроль срока. "
        "Из вашего процесса я зафиксировала следующие обязательные требования…»"
    )
    doc.add_paragraph("Назовите вслух 2–4 требования клиента его словами. Затем предложите только один конкретный следующий шаг.")

    add_table(
        doc,
        ("Ситуация", "Следующий шаг"),
        (
            ("Функциональность в целом подходит", "Отправить коммерческое предложение и согласовать тестовый месяц."),
            ("Нужно проверить нормы клиента", "Получить один обезличенный пример нормы и показать его в системе."),
            ("Есть обязательный пробел", "Описать сценарий, оценить возможность и только затем обсуждать сроки."),
            ("Решение принимает другой человек", "Назначить короткую повторную встречу с ответственным лицом."),
            ("Клиент пока сравнивает варианты", "Отправить краткое резюме преимуществ именно для его процесса."),
        ),
        widths=(2.15, 4.45),
    )

    doc.add_heading("Финальные вопросы", level=2)
    add_bullet(doc, "Какая из показанных возможностей для вас наиболее ценна?")
    add_bullet(doc, "Что обязательно должно быть решено до начала тестирования?")
    add_bullet(doc, "Кто ещё участвует в принятии решения?")
    add_bullet(doc, "Какой следующий шаг и дата будут удобны?")

    doc.add_heading("После встречи", level=2)
    add_number(doc, "В тот же день отправьте краткое резюме договорённостей.")
    add_number(doc, "Отдельно перечислите обязательные требования и открытые вопросы.")
    add_number(doc, "Приложите руководство пользователя и актуальное коммерческое предложение, если это согласовано.")
    add_number(doc, "Зафиксируйте ответственного, следующий шаг и конкретную дату связи.")

    doc.add_heading("Короткая формула презентации", level=2)
    doc.add_paragraph(
        "Сначала слушаем процесс клиента. Затем показываем проблему на рабочем экране, источник правила в норме, "
        "действие в выдаче и результат в отчёте. Завершаем не перечнем функций, а согласованным следующим шагом."
    )

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
